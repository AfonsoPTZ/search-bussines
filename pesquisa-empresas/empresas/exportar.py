"""Geração das planilhas (XLSX com abas de fontes e alternativas; CSV separado por ';' para Excel BR).

Valores não confirmados aparecem com o sufixo "(não confirmado)" e a célula fica laranja.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from empresas.lote import ler_registros
from empresas.pesquisa import CAMPOS, CAMPOS_MULTIPLOS, _CAMPOS_CADASTRAIS, _PRIORIDADE, _PRIORIDADE_CADASTRO_PRIMEIRO
from empresas.util import normalizar

NAO_CONFIRMADO = " (não confirmado)"

COLUNAS = [
    ("linha", "Linha"), ("entrada", "Entrada"), ("status", "Status"), ("confianca", "Confiança"),
    ("nao_confirmados", "Campos não confirmados"),
    ("nome_comercial", "Nome comercial"), ("categoria", "Segmento (A–M)"), ("razao_social", "Razão social"),
    ("cnpj", "CNPJ"), ("situacao_cadastral", "Situação cadastral"), ("data_situacao", "Data da situação"),
    ("atividade_principal", "Atividade principal (CNAE)"), ("natureza_juridica", "Natureza jurídica"),
    ("porte", "Porte"), ("matriz_filial", "Matriz/Filial"),
    ("telefones", "Telefones"), ("whatsapp", "WhatsApp"), ("emails", "E-mails"), ("site", "Site"),
    ("instagram", "Instagram"), ("seguidores_instagram", "Seguidores Instagram (aprox.)"),
    ("facebook", "Facebook"), ("linkedin", "LinkedIn"), ("outros_canais", "Outros canais"),
    ("horario_funcionamento", "Horário (como publicado)"), ("avaliacao", "Avaliação"),
    ("endereco_completo", "Endereço completo"), ("endereco", "Endereço"), ("bairro", "Bairro"),
    ("cidade", "Cidade"), ("uf", "UF"), ("cep", "CEP"), ("google_maps", "Google Maps"),
    ("telefone_cadastro", "Telefone (cadastro Receita)"), ("email_cadastro", "E-mail (cadastro Receita)"),
    ("fontes_principais", "Fontes principais"), ("divergencias", "Divergências"),
    ("nao_encontrados", "Não encontrados"), ("observacoes", "Observações"), ("data_pesquisa", "Data da pesquisa"),
]

NOMES_CAMPOS = dict(COLUNAS)

LEGENDA = [
    ("Status", ""),
    ("identificada", "Empresa identificada e CNPJ confirmado na base da Receita."),
    ("provavel", "CNPJ preenchido com o candidato mais plausível, mas NÃO confirmado. Veja Observações e a aba Alternativas."),
    ("parcial", "Empresa localizada (site/perfil/contatos), sem CNPJ."),
    ("nao_encontrada", "Nada encontrado (os candidatos descartados, se houver, estão na aba Alternativas)."),
    ("ambigua", "(lotes antigos) Mais de uma empresa possível; nenhuma escolhida."),
    ("erro", "Falha técnica; rode 'empresas retomar' para tentar de novo."),
    ("", ""),
    ("Confiança", ""),
    ("alta", "CNPJ informado, publicado no site/perfil ou confirmado por várias evidências (nome, cidade, domínio, telefone, endereço)."),
    ("media", "Boa correspondência, mas com divergência ou outras empresas ativas com o mesmo nome (franquias)."),
    ("baixa", "Correspondência fraca ou dados não confirmados; confira antes de usar."),
    ("", ""),
    ("Não confirmado", "Valor encontrado mas sem prova de que é desta empresa (ex.: CNPJ provável, perfil achado só pelo "
                       "nome na busca, site que bloqueou a leitura, segmento inferido por palavras). Aparece com o "
                       "sufixo '(não confirmado)', célula laranja e listado em 'Campos não confirmados'."),
    ("", ""),
    ("Origem das informações (aba Fontes)", ""),
    ("entrada", "Informado na lista."),
    ("cadastro", "Dados abertos da Receita Federal (OpenCNPJ/BrasilAPI/minhareceita/CNPJá)."),
    ("site", "Página do site oficial (URL exata na aba Fontes)."),
    ("instagram", "Perfil público do Instagram ou trecho dele exibido pelo buscador."),
    ("cep", "Consulta pública de CEP (BrasilAPI/ViaCEP) a partir do CEP encontrado."),
    ("busca", "Resultado de buscador (evidência mais fraca)."),
    ("gerado", "Montado pela ferramenta (link do Google Maps, segmento inferido por palavras-chave)."),
    ("revisao", "Correção feita na revisão (Claude ou manual), com fonte."),
    ("", ""),
    ("Seguidores", "Número que o buscador mostra para o perfil (arredondado pelo Instagram e pode estar alguns dias desatualizado)."),
    ("Horário", "Como publicado no site/perfil; pode ser o horário da loja ou do atendimento (veja a fonte)."),
    ("Avaliação", "Só quando o próprio site publica (dados estruturados) ou quando registrada na revisão. Notas do Google "
                  "não são coletadas automaticamente (exigiria API paga ou raspagem proibida)."),
    ("Campos vazios", "Nada encontrado com fonte. Nada é inventado."),
]

_CORES_STATUS = {"identificada": "C6EFCE", "provavel": "FCE4D6", "parcial": "FFEB9C", "ambigua": "F8CBAD",
                 "nao_encontrada": "D9D9D9", "erro": "FF9999"}
_CORES_CONF = {"alta": "C6EFCE", "media": "FFEB9C", "baixa": "F8CBAD"}
_COR_NAO_CONFIRMADO = "FFD8A8"
_LINKS = ("site", "google_maps", "instagram", "facebook", "linkedin")


def _ordenar(campo, itens):
    prior = _PRIORIDADE_CADASTRO_PRIMEIRO if campo in _CAMPOS_CADASTRAIS else _PRIORIDADE
    return sorted(itens, key=lambda x: (not x.get("confirmado", True), prior.get(x["origem"], 9)))


def celula(reg, campo) -> tuple[str, bool, str]:
    """(texto para exibir, tem valor não confirmado?, valor bruto do primeiro item — para links/números)."""
    itens = _ordenar(campo, reg["campos"].get(campo, []))
    if not itens:
        return "", False, ""
    if campo not in CAMPOS_MULTIPLOS:
        i = itens[0]
        confirmado = i.get("confirmado", True)
        texto = i["valor"] if confirmado or campo == "seguidores_instagram" else i["valor"] + NAO_CONFIRMADO
        return texto, not confirmado, i["valor"]
    vistos, partes, algum_nao = set(), [], False
    for i in itens:
        k = normalizar(i["valor"])
        if k in vistos:
            continue
        vistos.add(k)
        confirmado = i.get("confirmado", True)
        algum_nao |= not confirmado
        partes.append(i["valor"] if confirmado else i["valor"] + NAO_CONFIRMADO)
    return " | ".join(partes), algum_nao, itens[0]["valor"]


def _endereco_completo(reg) -> tuple[str, bool]:
    partes = {c: celula(reg, c) for c in ("endereco", "bairro", "cidade", "uf", "cep")}
    bruto = {c: v[2] for c, v in partes.items()}
    if not any(bruto.values()):
        return "", False
    local = "/".join(x for x in (bruto["cidade"], bruto["uf"]) if x)
    texto = " - ".join(x for x in (bruto["endereco"], bruto["bairro"]) if x)
    texto = ", ".join(x for x in (texto, local, f"CEP {bruto['cep']}" if bruto["cep"] else "") if x)
    nao = any(v[1] for v in partes.values())
    return (texto + NAO_CONFIRMADO if nao else texto), nao


def linha_planilha(reg) -> dict:
    """Valores de exibição da linha + chave '_nao_confirmados' (campos com valor não confirmado)."""
    d, nao = {}, set()
    for c in CAMPOS:
        texto, n, bruto = celula(reg, c)
        d[c] = texto
        d["_bruto_" + c] = bruto
        if n:
            nao.add(c)
    d["endereco_completo"], n = _endereco_completo(reg)
    if n:
        nao.add("endereco_completo")
    d.update(
        linha=reg["linha"], entrada=reg["entrada"], status=reg["status"], confianca=reg["confianca"],
        divergencias=" | ".join(reg["divergencias"]), observacoes=" | ".join(reg["observacoes"]),
        nao_encontrados=", ".join(reg["nao_encontrados"]), data_pesquisa=reg["data_pesquisa"],
        nao_confirmados=", ".join(NOMES_CAMPOS.get(c, c) for c in CAMPOS if c in nao),
    )
    d["_nao_confirmados"] = nao
    fontes = []
    for c in ("cnpj", "telefones", "emails", "site", "endereco"):
        for i in _ordenar(c, reg["campos"].get(c, [])):
            if i["fonte"] and i["fonte"] not in fontes:
                fontes.append(i["fonte"])
    d["fontes_principais"] = " | ".join(fontes[:5])
    return d


def _formatar_aba(ws, larguras=None):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="305496")
        cel.alignment = Alignment(vertical="center", wrap_text=True)
    for i, col in enumerate(ws.columns, 1):
        maior = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(i)].width = (larguras or {}).get(i, min(max(10, maior + 2), 60))


def _valor_excel(coluna, linha):
    """Seguidores vira número (ordenável no Excel); o resto fica texto."""
    if coluna == "seguidores_instagram":
        bruto = linha.get("_bruto_seguidores_instagram", "")
        return int(bruto) if str(bruto).isdigit() else None
    return linha.get(coluna, "")


def _csv_seguro(v):
    """No CSV, valor que começa com = ou @ ganha um apóstrofo para o Excel não executá-lo como fórmula."""
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "@") else v


def exportar(pasta: Path) -> dict[str, Path]:
    regs = [r for _, r in sorted(ler_registros(pasta).items())]
    linhas = [linha_planilha(r) for r in regs]

    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append([t for _, t in COLUNAS])
    idx = {c: i for i, (c, _) in enumerate(COLUNAS, 1)}
    laranja = PatternFill("solid", fgColor=_COR_NAO_CONFIRMADO)
    for n, l in enumerate(linhas, 2):
        ws.append([_valor_excel(c, l) for c, _ in COLUNAS])
        for cel in ws[n]:
            if cel.data_type == "f":  # valor raspado da web que começa com "=" é texto, não fórmula
                cel.data_type = "s"
        if cor := _CORES_STATUS.get(l["status"]):
            ws.cell(n, idx["status"]).fill = PatternFill("solid", fgColor=cor)
        if cor := _CORES_CONF.get(l["confianca"]):
            ws.cell(n, idx["confianca"]).fill = PatternFill("solid", fgColor=cor)
        for c in l["_nao_confirmados"]:
            if c in idx:
                cel = ws.cell(n, idx[c])
                cel.fill = laranja
                cel.font = Font(italic=True)
        for c in _LINKS:
            bruto = l.get("_bruto_" + c, "")
            if bruto.startswith("http"):
                cel = ws.cell(n, idx[c])
                cel.hyperlink = bruto
                cel.font = Font(color="0563C1", underline="single", italic=c in l["_nao_confirmados"])
    _formatar_aba(ws)

    wf = wb.create_sheet("Fontes")
    cab_fontes = ["Linha", "Entrada", "Campo", "Valor", "Confirmado", "Origem", "Fonte (URL)"]
    wf.append(cab_fontes)
    fontes_csv = []
    for r in regs:
        for campo in CAMPOS:
            for i in _ordenar(campo, r["campos"].get(campo, [])):
                row = [r["linha"], r["entrada"], campo, i["valor"], "sim" if i.get("confirmado", True) else "não",
                       i["origem"], i["fonte"]]
                wf.append(row)
                fontes_csv.append(row)
                if row[4] == "não":
                    wf.cell(wf.max_row, 5).fill = laranja
    _formatar_aba(wf)

    wc = wb.create_sheet("Alternativas")
    wc.append(["Linha", "Entrada", "Status", "Tipo", "Escolhido?", "Valor (CNPJ/site)", "Razão social", "Nome fantasia",
               "Cidade", "Situação", "Pontuação", "Motivos", "Encontrado em"])
    for r in regs:
        for c in r.get("candidatos_cnpj", []):
            wc.append([r["linha"], r["entrada"], r["status"], "CNPJ", c.get("escolhido", ""), c["cnpj"], c["razao_social"],
                       c["nome_fantasia"], c["cidade"], c["situacao"], c["pontuacao"], c["motivos"], c["encontrado_em"]])
        for a in r.get("alternativas", []):
            wc.append([r["linha"], r["entrada"], r["status"], a["tipo"], "não", a["valor"], "", "", "", "", "",
                       a["motivo"], a.get("fonte", "")])
    _formatar_aba(wc)

    wq = wb.create_sheet("Consultas")
    wq.append(["Linha", "Entrada", "Consulta / página lida"])
    for r in regs:
        for q in r.get("consultas", []):
            wq.append([r["linha"], r["entrada"], q])
    _formatar_aba(wq)

    wl = wb.create_sheet("Legenda")
    wl.append(["Termo", "Significado"])
    for linha in LEGENDA:
        wl.append(list(linha))
    wl.cell(next(n for n, l in enumerate(LEGENDA, 2) if l[0] == "Não confirmado"), 1).fill = laranja
    _formatar_aba(wl, {1: 34, 2: 110})

    arquivos = {"xlsx": pasta / "empresas.xlsx", "csv": pasta / "empresas.csv", "fontes": pasta / "fontes.csv"}
    try:
        wb.save(arquivos["xlsx"])
    except PermissionError:  # planilha aberta no Excel
        arquivos["xlsx"] = pasta / "empresas (nova).xlsx"
        wb.save(arquivos["xlsx"])
    with open(arquivos["csv"], "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([t for _, t in COLUNAS])
        for l in linhas:
            w.writerow(["" if (v := _valor_excel(c, l)) is None else _csv_seguro(v) for c, _ in COLUNAS])
    with open(arquivos["fontes"], "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(cab_fontes)
        w.writerows(fontes_csv)
    return arquivos
