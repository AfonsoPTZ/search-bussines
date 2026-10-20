"""Regressões dos achados da revisão de código (cada teste reproduz um cenário que falhava)."""

import json
import shutil

import pytest

from plataforma import acoes, base, db, exportar, sincronizar
from tests.conftest import Ambiente, perfil

R = Ambiente.registro


def _empresas():
    return [dict(r) for r in db.conexao().execute("SELECT * FROM empresas WHERE visivel=1 ORDER BY id")]


def test_pasta_movida_e_devolvida_preserva_unioes_e_vinculos(amb, tmp_path):
    pasta = amb.lote("2026-10-18_100000_a", [R(1, "Padaria Teste | Bauru/SP"), R(2, "@padaria_teste")])
    amb.lote("2026-10-18_100001_outro", [R(1, "Mercado Teste | Bauru/SP")])
    sincronizar.sincronizar()
    a, b = sorted(e["id"] for e in _empresas() if e["nome"] != "Mercado Teste")
    with db.transacao() as c:
        base.unir(c, a, b)
        base.decidir_vinculo(c, "padaria_teste", a, "confirmado", "é a mesma")
    fora = tmp_path / "fora"
    shutil.move(str(pasta), str(fora))
    sincronizar.sincronizar()
    assert [e["nome"] for e in _empresas()] == ["Mercado Teste"]
    shutil.move(str(fora), str(pasta))
    sincronizar.sincronizar()
    padaria = [e for e in _empresas() if e["nome"] != "Mercado Teste"]
    assert len(padaria) == 1 and padaria[0]["id"] == a and padaria[0]["n_pesquisas"] == 2
    v = db.conexao().execute("SELECT * FROM vinculos WHERE usuario='padaria_teste'").fetchone()
    assert v["empresa_id"] == a and v["manual"] == 1


def test_raiz_indisponivel_nao_apaga_nada(amb, monkeypatch, tmp_path):
    amb.lote("2026-10-18_100000_a", [R(1, "Padaria Teste | Bauru/SP")])
    sincronizar.sincronizar()
    from empresas import config as cfg_emp
    monkeypatch.setattr(cfg_emp, "RESULTADOS_DIR", tmp_path / "nao-existe")
    assert sincronizar.sincronizar().get("removidas", 0) == 0
    assert len(_empresas()) == 1


def test_mesma_entrada_com_sites_confirmados_diferentes_nao_une(amb):
    amb.lote("2026-10-18_100000_a", [R(1, "Funeraria Central", {"site": [("https://centralcuritiba.com.br/", "site", True)]})])
    amb.lote("2026-10-19_100000_b", [R(1, "Funeraria Central", {"site": [("https://centralrecife.com.br/", "site", True)]})])
    sincronizar.sincronizar()
    assert len(_empresas()) == 2
    assert "site/@/CNPJ confirmados diferentes" in db.conexao().execute("SELECT motivo FROM duplicidades").fetchone()[0]


def test_correcao_que_remove_o_identificador_desfaz_a_uniao_automatica(amb):
    amb.lote("2026-10-18_100000_a", [R(1, "Padaria X | Bauru/SP", {"instagram": [("https://www.instagram.com/padaria_x/", "site", True)]})])
    pb = amb.lote("2026-10-19_100000_b", [R(1, "Padaria X Centro | Bauru/SP", {"instagram": [("https://www.instagram.com/padaria_x/", "site", True)]})])
    sincronizar.sincronizar()
    assert len(_empresas()) == 1
    (pb / "registros.jsonl").write_text(json.dumps(R(1, "Padaria X Centro | Bauru/SP", revisado_em="2026-10-20T10:00:00")) + "\n", encoding="utf-8")
    sincronizar.sincronizar()
    assert len(_empresas()) == 2
    dono = db.conexao().execute("SELECT empresa_id FROM empresa_chaves WHERE chave='ig:padaria_x'").fetchone()[0]
    assert dono == db.conexao().execute("SELECT empresa_id FROM pesquisas WHERE id='emp:2026-10-18_100000_a#1'").fetchone()[0]


def test_uniao_manual_nao_e_desfeita_por_correcao(amb):
    amb.lote("2026-10-18_100000_a", [R(1, "Padaria Um | Bauru/SP"), R(2, "Padaria Dois | Bauru/SP")])
    sincronizar.sincronizar()
    a, b = [e["id"] for e in _empresas()]
    with db.transacao() as c:
        base.unir(c, a, b)
    sincronizar.sincronizar(forcar=True)
    assert len(_empresas()) == 1


def test_unir_mantem_a_rejeicao_manual(amb):
    amb.lote("2026-10-18_100000_a", [R(1, "Cemiterio A | Curitiba/PR", {"site": [("https://cemiterioa.com.br/", "site", True)]}),
                                     R(2, "Cemiterio A Ltda | Curitiba/PR")])
    amb.analise("origem", [perfil("perfil_x", "Cemitério", "cemitério parque", link="https://cemiterioa.com.br/")])
    sincronizar.sincronizar()
    ids = {e["nome"]: e["id"] for e in _empresas()}
    a, b = ids["Cemiterio A"], ids["Cemiterio A Ltda"]
    with db.transacao() as c:
        base.decidir_vinculo(c, "perfil_x", b, "rejeitado")
        base.unir(c, a, b)
    v = db.conexao().execute("SELECT * FROM vinculos WHERE usuario='perfil_x' AND empresa_id=?", (a,)).fetchone()
    assert v["estado"] == "rejeitado" and v["manual"] == 1


def test_registros_antigos_sem_confirmacao_e_ambiguos_nao_geram_chave_forte(amb):
    antigo = R(1, "Loja Um | Bauru/SP")
    antigo["campos"]["site"] = [{"valor": "https://shopping.com.br/", "fonte": "x", "origem": "busca"}]  # sem "confirmado"
    amb.lote("2026-10-18_100000_a", [antigo, R(2, "Loja Dois | Marília/SP", {"site": [("https://shopping.com.br/", "site", True)]}, status="ambigua")])
    sincronizar.sincronizar()
    assert len(_empresas()) == 2
    assert not db.conexao().execute("SELECT 1 FROM empresa_chaves WHERE chave LIKE 'dom:%'").fetchone()


def test_duplo_retomar_inicia_um_so_processo(amb):
    amb.lote("2026-10-20_100000_a", [R(1, "Padaria Teste | Bauru/SP")], status="interrompido")
    sincronizar.sincronizar()
    acoes.retomar("emp:2026-10-20_100000_a")
    with pytest.raises(acoes.AcaoInvalida, match="acabou de ser iniciada"):
        acoes.retomar("emp:2026-10-20_100000_a")
    assert len(amb.processos) == 1


def test_execucao_criada_durante_a_sincronia_nao_e_apagada(amb, monkeypatch):
    amb.lote("2026-10-20_100000_a", [R(1, "Padaria Teste | Bauru/SP")])
    original = sincronizar.pastas_claude

    def no_meio():  # uma requisição registra uma execução nova enquanto a sincronia percorre as pastas
        pasta = amb.lote("2026-10-20_100001_nova", [R(1, "Mercado Teste | Bauru/SP")])
        sincronizar.sincronizar_execucao(db.conexao(), "empresas", pasta, forcar=True)
        return original()

    monkeypatch.setattr(sincronizar, "pastas_claude", no_meio)
    sincronizar.sincronizar()
    assert db.conexao().execute("SELECT 1 FROM execucoes WHERE id='emp:2026-10-20_100001_nova'").fetchone()


def test_perfil_de_analise_apagada_e_reimportado_da_que_sobrou(amb):
    amb.analise("origem1", [perfil("teste_flores", "Floricultura Nova", "floricultura")], nome="2026-10-20_100001_um")
    p2 = amb.analise("origem2", [perfil("teste_flores", "Floricultura Velha", "floricultura")], nome="2026-10-20_100002_dois")
    linhas = [json.loads(l) for l in (p2 / "perfis.jsonl").read_text(encoding="utf-8").splitlines()]
    linhas[0]["em"] = "2026-10-21T00:00:00"  # a análise 2 é a mais recente
    (p2 / "perfis.jsonl").write_text("\n".join(json.dumps(l, ensure_ascii=False) for l in linhas) + "\n", encoding="utf-8")
    sincronizar.sincronizar()
    assert db.conexao().execute("SELECT nome FROM perfis").fetchone()[0] == "Floricultura Velha"
    shutil.rmtree(p2)
    sincronizar.sincronizar()
    sincronizar.sincronizar()
    assert db.conexao().execute("SELECT nome FROM perfis").fetchone()[0] == "Floricultura Nova"


@pytest.mark.parametrize("nome", ["emp:C:", "emp:.. ", "emp:...", "emp:2026-10-20_100000_a ", "emp:2026-10-20_100000_a.",
                                  "emp:../emp", "ig:..\\x"])
def test_nomes_de_execucao_invalidos_no_windows(amb, nome):
    amb.lote("2026-10-20_100000_a", [R(1, "Padaria Teste | Bauru/SP")])
    assert sincronizar.pasta_da_execucao(nome) is None
    assert sincronizar.pasta_da_execucao("emp:2026-10-20_100000_a") is not None


def test_valor_da_web_nao_vira_formula_na_planilha(amb):
    amb.lote("2026-10-20_100000_a", [R(1, "Loja Teste | Bauru/SP", {"horario_funcionamento": [('=HYPERLINK("http://x","clique")', "site", True)],
                                                                   "telefones": [("+55 14 3333-3333", "site", True)]})])
    sincronizar.sincronizar()
    assert exportar.texto_seguro('=HYPERLINK("x")').startswith("'=") and exportar.texto_seguro("+55 14 3333") == "+55 14 3333"
    from io import BytesIO

    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(exportar.xlsx_bytes({}))).active
    assert all(c.data_type != "f" for linha in ws.iter_rows() for c in linha)
